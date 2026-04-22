from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Pt

from ..profile.schema import Profile
from .tailor import TailoredResume


def render_docx(profile: Profile, tailored: TailoredResume, out_path: Path) -> Path:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    p = profile.personal
    h = doc.add_paragraph()
    run = h.add_run(p.full_name or "Your Name")
    run.bold = True
    run.font.size = Pt(18)

    contact_parts = [x for x in (p.email, p.phone, f"{p.city}, {p.state}".strip(", "), p.linkedin_url, p.github_url) if x]
    if contact_parts:
        doc.add_paragraph(" · ".join(contact_parts))

    if tailored.summary:
        doc.add_heading("Summary", level=1)
        doc.add_paragraph(tailored.summary)

    if tailored.skills:
        doc.add_heading("Skills", level=1)
        doc.add_paragraph(" · ".join(tailored.skills))

    if tailored.experience:
        doc.add_heading("Experience", level=1)
        for role in tailored.experience:
            header = doc.add_paragraph()
            r = header.add_run(f"{role.company} — {role.title}")
            r.bold = True
            dates = f"  {role.start} – {role.end or 'Present'}"
            if role.location:
                dates += f" · {role.location}"
            header.add_run(dates).italic = True
            for b in role.bullets:
                doc.add_paragraph(b, style="List Bullet")

    if tailored.projects:
        doc.add_heading("Projects", level=1)
        for proj in tailored.projects:
            p_para = doc.add_paragraph()
            p_para.add_run(proj.name).bold = True
            if proj.description:
                p_para.add_run(f" — {proj.description}")
            if proj.url:
                p_para.add_run(f" · {proj.url}")
            for b in proj.bullets:
                doc.add_paragraph(b, style="List Bullet")

    if tailored.education:
        doc.add_heading("Education", level=1)
        for ed in tailored.education:
            e_para = doc.add_paragraph()
            e_para.add_run(ed.institution).bold = True
            details = f" — {ed.degree}"
            if ed.field:
                details += f", {ed.field}"
            if ed.start or ed.end:
                details += f"  ({ed.start} – {ed.end})"
            e_para.add_run(details)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    return out_path
